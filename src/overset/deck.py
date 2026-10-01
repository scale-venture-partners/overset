"""Read a .pptx into what the rules need: geometry, text, fonts, colours.

Everything is in EMU (914,400 per inch), the unit the file itself uses, so a
rule never mixes units. Rendered positions (render.py) are converted into EMU
before any rule compares them with these frames.

python-pptx resolves a placeholder's position through its layout, but not its
font size or colour through the theme. Where a value is inherited rather than
stated, it is None here and the rules that need it skip that run rather than
guess.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn

EMU_PER_PT = 12_700

# Placeholders that are routinely left empty on purpose.
_QUIET_PLACEHOLDERS = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER}


@dataclass(frozen=True)
class Box:
    x: int
    y: int
    w: int
    h: int

    @property
    def right(self) -> int:
        return self.x + self.w

    @property
    def bottom(self) -> int:
        return self.y + self.h

    def contains(self, other: Box, tolerance: int = 0) -> bool:
        return (other.x >= self.x - tolerance and other.y >= self.y - tolerance
                and other.right <= self.right + tolerance and other.bottom <= self.bottom + tolerance)

    def intersection(self, other: Box) -> int:
        w = min(self.right, other.right) - max(self.x, other.x)
        h = min(self.bottom, other.bottom) - max(self.y, other.y)
        return w * h if w > 0 and h > 0 else 0

    @property
    def area(self) -> int:
        return max(self.w, 0) * max(self.h, 0)


@dataclass
class Run:
    text: str
    size_pt: float | None  # effective: stated size times any shrink-on-overflow scale
    font: str | None
    color: str | None  # "RRGGBB", only when stated as RGB


@dataclass
class Shape:
    name: str
    kind: str  # "text" | "picture" | "table" | "group" | "other"
    box: Box | None
    text: str = ""
    runs: list[Run] = field(default_factory=list)
    fill: str | None = None
    placeholder: str | None = None  # the placeholder type's name, if a placeholder
    quiet_placeholder: bool = False
    auto_grow: bool = False  # the frame resizes to its text, so text past it is not overflow
    image_aspect: float | None = None  # pictures: aspect of the visible (cropped) source
    from_layout: bool = False  # drawn by the layout or master, not the slide


@dataclass
class Slide:
    number: int  # 1-based
    layout: str
    shapes: list[Shape]

    @property
    def text_shapes(self) -> list[Shape]:
        return [s for s in self.shapes if s.text.strip()]

    @property
    def words(self) -> int:
        return sum(len(s.text.split()) for s in self.shapes if not s.from_layout)


@dataclass
class Deck:
    path: Path
    width: int
    height: int
    slides: list[Slide]

    @property
    def frame(self) -> Box:
        return Box(0, 0, self.width, self.height)


def _rgb(color_format) -> str | None:
    try:
        if color_format.type == MSO_COLOR_TYPE.RGB:
            return str(color_format.rgb)
    except (AttributeError, TypeError):
        pass
    return None


def _fill(shape) -> str | None:
    try:
        if shape.fill.type == MSO_FILL.SOLID:
            return _rgb(shape.fill.fore_color)
    except (AttributeError, TypeError, NotImplementedError):
        pass
    return None


def _font_scale(text_frame) -> float:
    """PowerPoint's shrink-text-on-overflow stores its scale on the body; the
    sizes on the runs are the pre-shrink ones."""
    norm = text_frame._txBody.bodyPr.find(qn("a:normAutofit"))
    if norm is not None and norm.get("fontScale"):
        return int(norm.get("fontScale")) / 100_000
    return 1.0


def _box(shape) -> Box | None:
    if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
        return None
    return Box(int(shape.left), int(shape.top), int(shape.width), int(shape.height))


def _runs(text_frame) -> list[Run]:
    scale = _font_scale(text_frame)
    runs = []
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            size = run.font.size.pt * scale if run.font.size is not None else None
            runs.append(Run(run.text, size, run.font.name, _rgb(run.font.color)))
    return runs


def _picture_aspect(picture) -> float | None:
    from PIL import Image

    try:
        with Image.open(io.BytesIO(picture.image.blob)) as image:
            iw, ih = image.size
    except Exception:  # an unreadable or vector image: no aspect to compare
        return None
    visible_w = iw * (1 - picture.crop_left - picture.crop_right)
    visible_h = ih * (1 - picture.crop_top - picture.crop_bottom)
    if visible_w <= 0 or visible_h <= 0:
        return None
    return visible_w / visible_h


def _shape(shape, from_layout=False) -> Shape:
    kind = "other"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        kind = "picture"
    elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        kind = "group"
    elif getattr(shape, "has_table", False) and shape.has_table:
        kind = "table"
    elif shape.has_text_frame:
        kind = "text"

    out = Shape(name=shape.name, kind=kind, box=_box(shape), fill=_fill(shape), from_layout=from_layout)
    if shape.is_placeholder:
        ph_type = shape.placeholder_format.type
        out.placeholder = getattr(ph_type, "name", str(ph_type))
        out.quiet_placeholder = ph_type in _QUIET_PLACEHOLDERS
    if shape.has_text_frame:
        out.text = shape.text_frame.text
        out.runs = _runs(shape.text_frame)
        out.auto_grow = shape.text_frame.auto_size == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    elif kind == "table":
        cells = [c.text_frame for row in shape.table.rows for c in row.cells]
        out.text = "\n".join(c.text for c in cells)
        out.runs = [r for c in cells for r in _runs(c)]
    elif kind == "group":
        out.text = "\n".join(_shape(child).text for child in shape.shapes)
    if kind == "picture":
        out.image_aspect = _picture_aspect(shape)
    return out


def load(path: str | Path) -> Deck:
    path = Path(path)
    prs = Presentation(str(path))
    slides = []
    for number, slide in enumerate(prs.slides, start=1):
        shapes = [_shape(s) for s in slide.shapes]
        # The layout and master draw text too -- a footer, a logo lockup -- and
        # a renderer's words have to be attributable to some frame.
        for inherited in (slide.slide_layout.shapes, slide.slide_layout.slide_master.shapes):
            shapes += [_shape(s, from_layout=True) for s in inherited if not s.is_placeholder]
        slides.append(Slide(number, slide.slide_layout.name, shapes))
    return Deck(path, int(prs.slide_width), int(prs.slide_height), slides)

