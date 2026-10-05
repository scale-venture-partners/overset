"""Rules that read only the .pptx."""

from overset.engine import Context
from overset.rules import load_rules
from overset.rules.base import REGISTRY
from overset.settings import Settings

load_rules()


def check(code, deck, settings=None):
    return REGISTRY[code].check(Context(deck, settings or Settings()))


def test_a_text_frame_past_the_edge_is_reported_and_one_inside_is_not(builder, make_deck):
    builder.text("Inside", x=1, y=1)
    builder.text("Past the top", x=1, y=-0.5)
    (f,) = check("OVS004", make_deck(builder))
    assert f.slide == 1 and f.snippet == "Past the top" and "extends past" in f.message


def test_pictures_may_bleed_off_the_edge(builder, make_deck):
    builder.picture(x=-1, y=-1, w=15, h=9)
    assert check("OVS004", make_deck(builder)) == []


def test_small_text_is_reported_at_its_effective_size(builder, make_deck):
    builder.text("Fine print", size=8)
    builder.text("Body", y=3, size=18)
    (f,) = check("OVS101", make_deck(builder))
    assert f.message == "text at 8.0pt, under 10pt" and f.snippet == "Fine print"


def test_shrink_to_fit_counts_against_the_floor(builder, make_deck):
    from pptx.oxml.ns import qn

    box = builder.text("Shrunk", size=14)
    body = box.text_frame._txBody.bodyPr
    for child in list(body):
        body.remove(child)
    norm = body.makeelement(qn("a:normAutofit"), {"fontScale": "62500"})
    body.append(norm)
    (f,) = check("OVS101", make_deck(builder))
    assert "8.8pt" in f.message, "14pt shown at 62.5% is 8.75pt"


def test_fonts_are_only_checked_once_a_brand_set_is_configured(builder, make_deck):
    builder.text("Comic", font="Comic Sans MS")
    builder.text("Brand", y=3, font="Inter")
    deck = make_deck(builder)
    assert check("OVS102", deck) == []
    (f,) = check("OVS102", deck, Settings(fonts=["Inter", "Lora"]))
    assert "Comic Sans MS" in f.message


def test_a_weight_of_a_brand_font_is_the_brand_font(builder, make_deck):
    builder.text("Medium", font="Inter Medium")
    builder.text("Bold", y=2, font="Inter-Bold")
    builder.text("Sans", y=3, font="Interstate")
    (f,) = check("OVS102", make_deck(builder), Settings(fonts=["Inter"]))
    assert "Interstate" in f.message


def test_theme_fonts_pass_the_brand_check(builder, make_deck):
    builder.text("Theme", font="+mn-lt")
    assert check("OVS102", make_deck(builder), Settings(fonts=["Inter"])) == []


def test_colours_off_the_palette_are_reported_within_a_tolerance(builder, make_deck):
    builder.text("Brand green", color="00C853")
    builder.text("Nearly green", y=2, color="02C955")
    builder.text("Purple", y=3, color="8E24AA")
    deck = make_deck(builder)
    assert check("OVS103", deck) == []
    (f,) = check("OVS103", deck, Settings(palette=["00C853", "#1F1F1B"]))
    assert f.message == "#8E24AA is not in the brand palette"


def test_a_solid_fill_off_the_palette_is_reported(builder, make_deck):
    box = builder.text("x", color="1F1F1B")
    box.fill.solid()
    from pptx.dml.color import RGBColor

    box.fill.fore_color.rgb = RGBColor.from_string("FF00FF")
    (f,) = check("OVS103", make_deck(builder), Settings(palette=["1F1F1B"]))
    assert "FF00FF" in f.message and "fill of" in f.snippet


def test_an_empty_title_placeholder_is_reported(builder, make_deck):
    builder.new_slide(layout=1)  # title and content, both left empty
    found = check("OVS201", make_deck(builder))
    assert {f.message.split()[1] for f in found} == {"title", "object"}


def test_a_filled_placeholder_is_not(builder, make_deck):
    slide = builder.new_slide(layout=5)  # title only
    slide.shapes.title.text = "A claim, not a label"
    assert check("OVS201", make_deck(builder)) == []


def test_a_stretched_picture_is_reported_and_a_proportional_one_is_not(builder, make_deck):
    builder.picture(size=(400, 200), w=4, h=2)
    builder.picture(size=(400, 200), y=4, w=4, h=4)
    (f,) = check("OVS202", make_deck(builder))
    assert "stretched 50%" in f.message, "a 2:1 image shown 1:1"


def test_word_count_has_a_ceiling(builder, make_deck):
    builder.text(" ".join(["word"] * 120), h=5)
    (f,) = check("OVS203", make_deck(builder))
    assert f.message == "120 words, over 90"
    assert check("OVS203", make_deck(builder), Settings(max_words=200)) == []


def test_repeated_layouts_are_reported_when_asked_for(builder, make_deck):
    builder.new_slide(layout=6)
    builder.new_slide(layout=6)
    builder.new_slide(layout=5)
    (f,) = check("OVS301", make_deck(builder))
    assert f.slide == 2 and "same layout as slide 1" in f.message
