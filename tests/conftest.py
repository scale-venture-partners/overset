"""Decks built in code, so each test states exactly the defect it is about."""

from __future__ import annotations

import io
import shutil
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from overset.deck import Box, load
from overset.render import Page, Render, Word
from overset.settings import Settings

W, H = Inches(13.333), Inches(7.5)

needs_libreoffice = pytest.mark.skipif(
    not (shutil.which("soffice") or shutil.which("libreoffice")) or not shutil.which("pdftotext"),
    reason="needs LibreOffice and poppler",
)


class DeckBuilder:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.slide = None

    def new_slide(self, layout=6):
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[layout])
        return self.slide

    def text(self, text, x=1, y=1, w=6, h=1, size=24, color=None, font=None, grow=False):
        if self.slide is None:
            self.new_slide()
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        from pptx.enum.text import MSO_AUTO_SIZE

        tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT if grow else MSO_AUTO_SIZE.NONE
        tf.text = text
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if font:
            run.font.name = font
        return box

    def picture(self, size=(400, 200), x=1, y=1, w=4, h=2):
        if self.slide is None:
            self.new_slide()
        buf = io.BytesIO()
        Image.new("RGB", size, "red").save(buf, format="PNG")
        buf.seek(0)
        return self.slide.shapes.add_picture(buf, Inches(x), Inches(y), Inches(w), Inches(h))

    def save(self, path: Path) -> Path:
        self.prs.save(str(path))
        return path


@pytest.fixture
def builder():
    return DeckBuilder()


@pytest.fixture
def make_deck(tmp_path):
    def make(b: DeckBuilder, name="deck.pptx"):
        return load(b.save(tmp_path / name))

    return make


@pytest.fixture
def settings():
    return Settings()


def word(text, x, y, w, h, block=0, line=0) -> Word:
    """A rendered word, positioned in inches."""
    return Word(text, Box(int(Inches(x)), int(Inches(y)), int(Inches(w)), int(Inches(h))), block, line)


@pytest.fixture
def fake_render(tmp_path):
    """A Render with chosen words, over a plain image per page."""

    def make(pages_words, color="white", dpi=40, draw=None):
        pages = []
        for n, words in enumerate(pages_words, start=1):
            image = Image.new("RGB", (int(13.333 * dpi), int(7.5 * dpi)), color)
            if draw:
                draw(n, image, dpi)
            path = tmp_path / f"slide-{n}.png"
            image.save(path)
            pages.append(Page(n, words, path, dpi))
        return Render(pages, tmp_path)

    return make
